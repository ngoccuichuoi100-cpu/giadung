import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Page Setup - Margins 1 inch
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Helper for Colors
COLOR_PRIMARY = RGBColor(255, 104, 2)     # Orange #FF6802
COLOR_DARK = RGBColor(2, 2, 2)           # Onyx #020202
COLOR_TEXT = RGBColor(51, 51, 51)        # Slate #333333
COLOR_GRAY = RGBColor(100, 116, 139)     # Slate-500 #64748B

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

# Title
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_sub = title_p.add_run("BÁO CÁO KẾ HOẠCH & BẢN ĐỒ ĐỊNH HƯỚNG\n")
run_sub.font.size = Pt(11)
run_sub.font.bold = True
run_sub.font.color.rgb = COLOR_PRIMARY

run_title = title_p.add_run("LỘ TRÌNH THỰC HIỆN DỰ ÁN GIADỤNGPRO (ROADMAP)")
run_title.font.name = "Arial"
run_title.font.size = Pt(20)
run_title.font.bold = True
run_title.font.color.rgb = COLOR_DARK

doc.add_paragraph()

# Overview Box
p_meta = doc.add_paragraph()
p_meta.paragraph_format.space_after = Pt(12)
run_m = p_meta.add_run("📌 THÔNG TIN TỔNG QUAN DỰ ÁN:\n")
run_m.font.bold = True
run_m.font.size = Pt(11)
run_m.font.color.rgb = COLOR_PRIMARY

p_m_content = p_meta.add_run(
    "• Tên dự án: Hệ thống Thương mại Điện tử Thiết bị Gia dụng Cao cấp (GiaDụngPro)\n"
    "• Công nghệ sử dụng: Laravel 12 (PHP 8.3), Tailwind CSS v4, JavaScript ES6+, SQLite/MySQL\n"
    "• Mục tiêu: Xây dựng trải nghiệm mua sắm trực tuyến vượt trội với hiệu ứng 3D, Quick View và Gamification tích điểm."
)
p_m_content.font.size = Pt(10.5)
p_m_content.font.color.rgb = COLOR_TEXT

doc.add_paragraph()

def add_heading(text, level=1):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    run = h.add_run(text)
    run.font.name = "Arial"
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = COLOR_PRIMARY
    else:
        run.font.size = Pt(12)
        run.font.color.rgb = COLOR_DARK
    return h

# Section 1
add_heading("GIAI ĐOẠN 1: CHUẨN HÓA LUỒNG TRẢI NGHIỆM KHÁCH HÀNG (USER JOURNEY)", 1)

table1 = doc.add_table(rows=1, cols=3)
table1.alignment = WD_TABLE_ALIGNMENT.CENTER
table1.autofit = False

hdr_cells = table1.rows[0].cells
hdr_titles = ["Bước / Chức năng", "Mô tả chi tiết trải nghiệm", "Trạng thái"]
for i, title in enumerate(hdr_titles):
    hdr_cells[i].text = title
    set_cell_background(hdr_cells[i], "FF6802")
    for paragraph in hdr_cells[i].paragraphs:
        for run in paragraph.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(10)

data_g1 = [
    ("1.1 Khám phá & Tìm kiếm", "Trang chủ hiển thị Banner 3D Carousel, Thương hiệu Coverflow, Danh mục nổi bật và Thanh tìm kiếm thông minh.", "✅ Hoàn thành"),
    ("1.2 Tương tác Sản phẩm 3D", "Mở Quick View Modal trực tiếp trên thẻ sản phẩm đúng sơ đồ minh họa. Xem chi tiết với 3D Lightbox Zoom (Zoom, xoay 180°, drag).", "✅ Hoàn thành"),
    ("1.3 Tùy chọn Biến thể & Bảo hành", "Lựa chọn dải màu Hex Swatches, thông số kỹ thuật và đăng ký các gói bảo hành vàng nâng cao.", "✅ Hoàn thành"),
    ("1.4 Giỏ hàng & Đặt hàng", "Thêm vào giỏ (xử lý Session + DB), áp dụng Coupon giảm giá, dùng điểm thưởng tích lũy và thực hiện Checkout.", "✅ Hoàn thành"),
    ("1.5 Lịch sử & Đánh giá", "Quản lý đơn hàng cá nhân, hủy đơn hoàn tồn kho, viết nhận xét sản phẩm kèm bình chọn đánh giá hữu ích.", "✅ Hoàn thành"),
]

for item, desc, status in data_g1:
    row_cells = table1.add_row().cells
    row_cells[0].text = item
    row_cells[1].text = desc
    row_cells[2].text = status
    for i in range(3):
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                run.font.color.rgb = COLOR_TEXT
                if i == 2:
                    run.font.bold = True

doc.add_paragraph()

# Section 2
add_heading("GIAI ĐOẠN 2: CHUẨN HÓA LUỒNG QUẢN TRỊ (ADMIN MANAGEMENT FLOW)", 1)

table2 = doc.add_table(rows=1, cols=3)
table2.alignment = WD_TABLE_ALIGNMENT.CENTER
table2.autofit = False

hdr_cells2 = table2.rows[0].cells
for i, title in enumerate(hdr_titles):
    hdr_cells2[i].text = title
    set_cell_background(hdr_cells2[i], "020202")
    for paragraph in hdr_cells2[i].paragraphs:
        for run in paragraph.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(10)

data_g2 = [
    ("2.1 Đăng nhập & Phân quyền", "Đăng nhập Admin với username/email 'admin' và mật khẩu '123'. Kiểm soát quyền truy cập hệ thống.", "✅ Hoàn thành"),
    ("2.2 Quản lý Đơn hàng", "Duyệt đơn hàng, thay đổi trạng thái (Chờ duyệt -> Đang giao -> Đã giao / Hủy đơn tự động hoàn kho).", "✅ Hoàn thành"),
    ("2.3 Quản lý Kho & Sản phẩm", "Thêm/Sửa/Xóa sản phẩm, danh mục, thương hiệu, biến thể màu sắc và cập nhật số lượng tồn kho.", "✅ Hoàn thành"),
    ("2.4 Quản lý Khuyến mãi & Gamification", "Tạo mã giảm giá Coupon, cấu hình Vòng quay may mắn (Lucky Wheel) và Quản lý điểm thưởng check-in.", "✅ Hoàn thành"),
]

for item, desc, status in data_g2:
    row_cells = table2.add_row().cells
    row_cells[0].text = item
    row_cells[1].text = desc
    row_cells[2].text = status
    for i in range(3):
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                run.font.color.rgb = COLOR_TEXT
                if i == 2:
                    run.font.bold = True

doc.add_paragraph()

# Section 3
add_heading("GIAI ĐOẠN 3: ĐÓNG GÓI, TỐI ƯU & TRÌNH DIỄN (SHOWCASE READY)", 1)

p_g3 = doc.add_paragraph()
p_g3_content = (
    "1. Kiểm thử Tối ưu hóa Giao diện (Responsive Test):\n"
    "   - Kiểm tra hiển thị chuẩn trên Điện thoại di động (Mobile) & Máy tính bảng (Tablet).\n"
    "   - Đảm bảo các khung Modal Quick View và Lightbox 3D hiển thị mượt mà không tràn viền.\n\n"
    "2. Kịch bản Thuyết minh / Bảo vệ Đồ án (Demo Checklist):\n"
    "   - Bước 1: Giới thiệu Trang chủ với Hero Banner 3D Carousel & Thương hiệu Coverflow.\n"
    "   - Bước 2: Trình diễn Quick View mở rộng thông tin thẻ sản phẩm trực tiếp.\n"
    "   - Bước 3: Trình diễn 3D Lightbox Zoom ảnh sản phẩm sắc nét.\n"
    "   - Bước 4: Đặt hàng thử nghiệm & Áp dụng Voucher / Điểm thưởng.\n"
    "   - Bước 5: Đăng nhập Admin duyệt đơn hàng & xem thống kê."
)
p_g3_run = p_g3.add_run(p_g3_content)
p_g3_run.font.size = Pt(10)
p_g3_run.font.color.rgb = COLOR_TEXT

doc.add_paragraph()

# Footer note
p_foot = doc.add_paragraph()
p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r_foot = p_foot.add_run("Tài liệu được tạo tự động bởi Antigravity AI Assistant • GiaDụngPro 2026")
r_foot.font.italic = True
r_foot.font.size = Pt(9)
r_foot.font.color.rgb = COLOR_GRAY

doc.save("c:/laragon/www/giadung/LO_TRINH_DU_AN_GIADUNGPRO.docx")
print("Successfully generated LO_TRINH_DU_AN_GIADUNGPRO.docx!")
